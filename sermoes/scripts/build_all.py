"""Monta a base unificada de sermões (planilha + JSON do dashboard)."""
import json, re, os, sys
from datetime import datetime, timedelta
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import biblia, meta
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

SRC = '/home/user/teste/sermoes/sermoes_fdqc_preenchida.xlsx'
OUT = '/home/user/teste/sermoes/sermoes_fdqc_completa.xlsx'
OUT_JSON = 'dashboard_data.json'

rest = json.load(open('rest.json'))
res5 = {}  # texto base dos episódios do podcast sem referência no título (k -> (texto, codigo))
if os.path.exists('res5.tsv'):
    for l in open('res5.tsv'):
        if l.strip():
            k, t, c = l.rstrip('\n').split('\t')
            res5[int(k)] = (t, c)
DUP = json.load(open('dups.json')) if os.path.exists('dups.json') else {}  # k(rest) -> n(lista)

CONF = {'A': 'Confirmado', 'M': 'Provável', 'B': 'Incerto', 'T': 'Temático', 'P': 'Panorama', 'N': 'Não identificado'}


def norm_status(s):
    s = s or ''
    if s.startswith('Confirmado') or 'texto lido: confirmado' in s: return 'Confirmado'
    if s.startswith('Provável') or s.startswith('Sugestão'): return 'Provável'
    if s.startswith('Incerto'): return 'Incerto'
    if s.startswith('Temático'): return 'Temático'
    if s.startswith('Panorama'): return 'Panorama'
    return 'Não identificado'


def fonte_de(status, nlm):
    if 'NotebookLM' in (status or '') and 'transcrição' in (status or ''): return 'Transcrição + NotebookLM'
    if nlm: return 'NotebookLM'
    return 'Transcrição do áudio'


def ref_titulo(t):
    cands = re.findall(r'\(([^()]*)\)', t) + re.findall(r'\[([^\]]*)\]', t) + re.findall(r'//\s*([^-]+?)\s+-', t) + [t]
    for c in cands:
        c2 = re.split(r'\s*-?\s*S[ée]rie\b|\s+-\s+S[ÉE]RIE', c)[0]
        if not biblia.parse(c2) and 'Série' in c:
            c2 = re.sub(r'^.*?S[ée]rie [^-]+-\s*', '', c)
        if c is t:
            m = re.search(r'((?:\d\s*|I\s+|1ª\s*)?(?:Salmos?|Romanos|Filipenses|Efésios|Gálatas|Lucas|Mateus|Eclesiastes|Apocalipse|Atos|Coríntios)\s+[\d:.;\-– ]+\d)', t)
            c2 = m.group(1) if m else ''
        c2 = re.split(r'\s+-\s+(?=[A-Za-zÀ-ú])', c2)[0]
        if biblia.parse(c2):
            return c2.strip(' -()')
    return ''


def limpa_titulo(t):
    t = re.sub(r'^•\s*|\s*•\s*$', '', t).replace(' • ', ' - ')
    return t.strip()


def tipo_de(titulo, serie, grupo=''):
    tl = titulo.lower()
    if 'especial de natal' in tl: return 'Especial de Natal'
    if serie in ('Podcast Ser Família', 'Vida Comum (podcast)') or 'podcast' in tl: return 'Podcast / conversa'
    if serie in ('Teologue', 'Conferência (RE)Pensando a Igreja', 'Com Todos os Santos') or 'encontro de mulheres' in tl or 'belmonte' in tl or 'guilherme de carvalho' in tl:
        return 'Palestra / conferência'
    if 'convidados especiais' in tl: return 'Outro'
    if serie == 'Ensino das Escrituras' or 'encontro da família' in tl: return 'Ensino / encontro'
    return 'Sermão'


def ano_aprox(p):
    m = re.search(r'(20\d\d)', p or '')
    return int(m.group(1)) if m else None


def domingo_antes(d):
    """Domingo anterior (ou igual) à data de publicação."""
    dt = datetime.strptime(d, '%d/%m/%Y')
    return (dt - timedelta(days=(dt.weekday() + 1) % 7)).strftime('%d/%m/%Y')


recs = []
wb = load_workbook(SRC)
ws = wb['Sermões']
dup_by_n = {v: int(k) for k, v in DUP.items()}
for row in ws.iter_rows(min_row=2, values_only=True):
    n, tit, preg, serie_pdf, grupo, data, periodo, texto, status, evid, pod, yt, nlm = (list(row) + [None] * 13)[:13]
    extra = ''
    if grupo and grupo.startswith('Séries'):
        extra = re.sub(r'\s*\(.*\)', '', serie_pdf or '')
    if serie_pdf in ('Encontro de Mulheres', 'Podcast Ser Família'):
        extra = serie_pdf
    if (serie_pdf or '').startswith('Igreja (primeiras'):
        extra = 'Igreja (primeiras'
    if serie_pdf == 'A Vida no Reino (~2023, provavelmente Mateus 5-7: confirmar)':
        extra = 'A Vida no Reino'
    serie = meta.serie(tit, extra)
    if extra == 'Encontro de Mulheres': serie = 'Encontro de Mulheres'
    p = meta.pregador(tit) or preg or ''
    if not p and 'Especial de Natal' in tit: p = 'Vários (especial)'
    k = dup_by_n.get(n)
    if not data and k is not None:
        data = rest[k]['d']
        pod = rest[k]['link']
    st = norm_status(status)
    if (status or '').startswith('Não se aplica') and st == 'Não identificado':
        st = 'Temático' if (not texto or 'vários' in texto) else 'Provável'
    if (status or '').startswith('Sem áudio'): st = 'Não identificado'
    recs.append(dict(
        id=f'Y{n:03d}', titulo=limpa_titulo(tit), pregador=p, data=data or '',
        domingo=domingo_antes(data) if data else '', ano=int(data[-4:]) if data else ano_aprox(periodo),
        data_aprox=not data, periodo=periodo or '', texto=texto or '', confianca=st,
        fonte=('Título do vídeo' if False else fonte_de(status, nlm)) if texto else '',
        serie=serie, tipo=tipo_de(tit, serie, grupo), youtube=yt or '', podcast=pod or '',
        evidencia=evid or '', notebooklm=nlm or ''))

for k, r in enumerate(rest):
    if str(k) in DUP:
        continue
    t = r['t']
    if t.startswith('Vida Comum') or 'IHOP' in t:
        pass
    serie = meta.serie(t)
    texto, st, fonte = '', 'Não identificado', ''
    rt = ref_titulo(t) if r['rt'] else ''
    if k in res5:
        texto, c = res5[k]
        st, fonte = CONF[c], 'Transcrição do áudio'
    elif rt and not re.search(r'Romanos (8-16|1-7)$', rt):
        texto, st, fonte = rt, 'Confirmado', 'Título do episódio'
    elif rt:
        texto, st, fonte = '', 'Não identificado', ''
    if not texto and r.get('rd') and k not in res5:
        texto, st, fonte = '; '.join(r['rd'][:2]), 'Provável', 'Descrição do episódio'
    recs.append(dict(
        id=f'P{k:03d}', titulo=t.strip(), pregador=meta.pregador(t), data=r['d'], domingo=domingo_antes(r['d']),
        ano=int(r['d'][-4:]), data_aprox=False, periodo='', texto=texto, confianca=st, fonte=fonte,
        serie=serie, tipo=tipo_de(t, serie), youtube='', podcast=r['link'], evidencia='', notebooklm=''))

# referências e livros
for x in recs:
    refs = biblia.parse(x['texto']) if x['confianca'] not in ('Não identificado',) else []
    x['refs'] = [[a, sorted(c)] for a, c in refs]
    x['livros'] = ', '.join(biblia.INFO[a]['nome'] for a, _ in refs)
    x['capitulos'] = '; '.join(biblia.fmt(a, c) for a, c in refs)
    ts = {biblia.INFO[a]['test'] for a, _ in refs}
    x['testamento'] = 'AT e NT' if len(ts) == 2 else ('Antigo Testamento' if ts == {'AT'} else 'Novo Testamento' if ts else '')
    x['link'] = x['youtube'] or x['podcast'] or ('https://www.youtube.com/@familiadosquecreem/search?query=' + re.sub(r'\s+', '+', re.split(r' - | \(|\[', x['titulo'])[0].strip()))
    x['em_serie'] = 'Sim' if x['serie'] else 'Não'

recs.sort(key=lambda x: (x['ano'] or 0, datetime.strptime(x['data'], '%d/%m/%Y') if x['data'] else datetime(x['ano'] or 1900, 1, 1)), reverse=True)

json.dump(dict(sermoes=recs, livros=[dict(ab=a, nome=n, test=t, caps=c, genero=biblia.GENERO[a]) for a, n, t, c, _ in biblia.LIVROS]),
          open(OUT_JSON, 'w'), ensure_ascii=False)

# ---------- planilha ----------
wb2 = Workbook()
s = wb2.active
s.title = 'Sermões'
cols = [('Data', 12, 'data'), ('Domingo provável', 12, 'domingo'), ('Ano', 7, 'ano'), ('Título', 60, 'titulo'),
        ('Pregador', 22, 'pregador'), ('Texto base', 30, 'texto'), ('Livro(s)', 20, 'livros'), ('Capítulos', 18, 'capitulos'),
        ('Testamento', 16, 'testamento'), ('Faz parte de série?', 10, 'em_serie'), ('Série', 32, 'serie'), ('Tipo', 18, 'tipo'),
        ('Link (abrir sermão)', 42, 'link'), ('Link podcast', 40, 'podcast'), ('Confiança do texto', 14, 'confianca'),
        ('Fonte do texto', 22, 'fonte'), ('Trecho da transcrição', 60, 'evidencia'), ('NotebookLM', 40, 'notebooklm'), ('ID', 7, 'id')]
s.append([c[0] for c in cols])
for x in recs:
    vals = []
    for _, _, k in cols:
        v = x[k]
        if k == 'data' and not v and x['periodo']:
            v = x['periodo'] + ' (aprox.)'
        vals.append(v if v != '' else None)
    s.append(vals)
    r = s.max_row
    for k in ('link', 'podcast'):
        ci = [c[2] for c in cols].index(k) + 1
        if x[k]:
            cell = s.cell(r, ci); cell.hyperlink = x[k]; cell.style = 'Hyperlink'
hdr = PatternFill('solid', fgColor='1F3A5F')
for c in s[1]:
    c.font = Font(bold=True, color='FFFFFF'); c.fill = hdr; c.alignment = Alignment(wrap_text=True, vertical='center')
for i, (_, w, _) in enumerate(cols, 1):
    s.column_dimensions[s.cell(1, i).column_letter].width = w
s.freeze_panes = 'E2'
s.auto_filter.ref = f'A1:{s.cell(1, len(cols)).column_letter}{s.max_row}'
s.row_dimensions[1].height = 32

# referências (uma linha por livro/capítulo) para tabelas dinâmicas
r2 = wb2.create_sheet('Referências')
r2.append(['ID', 'Data', 'Ano', 'Título', 'Pregador', 'Série', 'Livro', 'Testamento', 'Gênero', 'Capítulo', 'Confiança'])
for x in recs:
    for a, caps in x['refs']:
        for c in (caps or [None]):
            r2.append([x['id'], x['data'] or None, x['ano'], x['titulo'], x['pregador'], x['serie'] or None, biblia.INFO[a]['nome'],
                       biblia.INFO[a]['test'], biblia.GENERO[a], c, x['confianca']])
for c in r2[1]:
    c.font = Font(bold=True, color='FFFFFF'); c.fill = hdr
for col, w in zip('ABCDEFGHIJK', [7, 12, 7, 50, 22, 30, 18, 6, 18, 9, 13]):
    r2.column_dimensions[col].width = w
r2.auto_filter.ref = f'A1:K{r2.max_row}'
r2.freeze_panes = 'A2'

# cobertura por livro
cv = wb2.create_sheet('Cobertura por livro')
cv.append(['Livro', 'Testamento', 'Gênero', 'Capítulos no livro', 'Capítulos pregados', '% coberto', 'Nº de sermões', 'Capítulos nunca pregados'])
for a, n, t, caps, _ in biblia.LIVROS:
    feitos, ns = set(), 0
    for x in recs:
        for b, cs in x['refs']:
            if b == a:
                feitos.update(cs); ns += 1
    falta = [c for c in range(1, caps + 1) if c not in feitos]
    cv.append([n, t, biblia.GENERO[a], caps, len(feitos), round(len(feitos) / caps, 3), ns,
               biblia.fmt('', set(falta)).strip() if falta and len(falta) < caps else ('todos' if falta else '')])
    cv.cell(cv.max_row, 6).number_format = '0%'
for c in cv[1]:
    c.font = Font(bold=True, color='FFFFFF'); c.fill = hdr; c.alignment = Alignment(wrap_text=True)
for col, w in zip('ABCDEFGH', [20, 10, 18, 11, 11, 10, 11, 60]):
    cv.column_dimensions[col].width = w
cv.freeze_panes = 'A2'

info = wb2.create_sheet('Sobre')
for l in ['Sermões da Família dos Que Creem — base completa', '',
          f'{len(recs)} gravações: os 240 vídeos da lista do YouTube + os episódios do podcast (SoundCloud) que não estavam na lista.',
          '• Data: publicação no podcast (geralmente 2 a 4 dias após o culto). "Domingo provável" = domingo anterior à publicação.',
          '  Vídeos sem áudio no podcast têm só o período aproximado.',
          '• Link (abrir sermão): YouTube quando conhecido; senão o episódio do podcast.',
          '• Confiança do texto: Confirmado (anunciado no sermão ou no título) · Provável · Incerto · Temático (sem texto único) · Panorama (introdução a um livro).',
          '• Fonte do texto: Título do episódio · Transcrição do áudio (10 primeiros minutos) · NotebookLM · Descrição do episódio.',
          '• Aba "Referências": uma linha por livro/capítulo, para tabelas dinâmicas. Aba "Cobertura por livro": o que já foi e o que nunca foi pregado.']:
    info.append([l])
info.column_dimensions['A'].width = 140
info['A1'].font = Font(bold=True, size=13)
wb2.save(OUT)

from collections import Counter
print(len(recs), Counter(x['confianca'] for x in recs), Counter(x['tipo'] for x in recs))
print('sem pregador:', [x['titulo'][:50] for x in recs if not x['pregador']])
print('séries:', len({x['serie'] for x in recs if x['serie']}))
