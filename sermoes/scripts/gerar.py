"""Gera a planilha e o dashboard a partir de dados/sermoes.json.

Uso: python sermoes/scripts/gerar.py
"""
import json, os, re, sys
from datetime import datetime
AQUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, AQUI)
import biblia
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

RAIZ = os.path.dirname(AQUI)
BASE = os.path.join(RAIZ, 'dados', 'sermoes.json')
OUT = os.path.join(RAIZ, 'sermoes_fdqc_completa.xlsx')
DASH = os.path.join(RAIZ, 'dashboard.html')


def derivar(x):
    """Recalcula livros, capítulos, testamento, link e domingo a partir do texto base."""
    refs = biblia.parse(x['texto']) if x['confianca'] != 'Não identificado' else []
    x['refs'] = [[a, sorted(c)] for a, c in refs]
    x['livros'] = ', '.join(biblia.INFO[a]['nome'] for a, _ in refs)
    x['capitulos'] = '; '.join(biblia.fmt(a, c) for a, c in refs)
    ts = {biblia.INFO[a]['test'] for a, _ in refs}
    x['testamento'] = 'AT e NT' if len(ts) == 2 else ('Antigo Testamento' if ts == {'AT'} else 'Novo Testamento' if ts else '')
    x['link'] = x.get('youtube') or x.get('podcast') or x.get('link', '')
    x['em_serie'] = 'Sim' if x.get('serie') else 'Não'
    if x.get('data'):
        dt = datetime.strptime(x['data'], '%d/%m/%Y')
        from datetime import timedelta
        x['domingo'] = (dt - timedelta(days=(dt.weekday() + 1) % 7)).strftime('%d/%m/%Y')
        x['ano'] = dt.year
    return x


dados = json.load(open(BASE))
recs = [derivar(x) for x in dados['sermoes']]
recs.sort(key=lambda x: (x['ano'] or 0, datetime.strptime(x['data'], '%d/%m/%Y') if x['data'] else datetime(x['ano'] or 1900, 1, 1)), reverse=True)
dados['sermoes'] = recs
json.dump(dados, open(BASE, 'w'), ensure_ascii=False, indent=0)

# ---------- planilha ----------
wb2 = Workbook()
s = wb2.active
s.title = 'Sermões'
cols = [('Data', 12, 'data'), ('Domingo provável', 12, 'domingo'), ('Ano', 7, 'ano'), ('Título', 60, 'titulo'),
        ('Pregador', 22, 'pregador'), ('Texto base', 30, 'texto'), ('Livro(s)', 20, 'livros'), ('Capítulos', 18, 'capitulos'),
        ('Testamento', 16, 'testamento'), ('Faz parte de série?', 10, 'em_serie'), ('Série', 32, 'serie'), ('Tipo', 18, 'tipo'),
        ('Link (abrir sermão)', 42, 'link'), ('Link podcast', 40, 'podcast'), ('Confiança do texto', 14, 'confianca'),
        ('Fonte do texto', 22, 'fonte'), ('Trecho da transcrição', 60, 'evidencia'), ('NotebookLM', 40, 'notebooklm'), ('ID', 7, 'id')]
s.append([c[0] for c in cols])
for x in recs:
    vals = []
    for _, _, k in cols:
        v = x[k]
        if k == 'data' and not v and x['periodo']:
            v = x['periodo'] + ' (aprox.)'
        vals.append(v if v != '' else None)
    s.append(vals)
    r = s.max_row
    for k in ('link', 'podcast'):
        ci = [c[2] for c in cols].index(k) + 1
        if x[k]:
            cell = s.cell(r, ci); cell.hyperlink = x[k]; cell.style = 'Hyperlink'
hdr = PatternFill('solid', fgColor='1F3A5F')
for c in s[1]:
    c.font = Font(bold=True, color='FFFFFF'); c.fill = hdr; c.alignment = Alignment(wrap_text=True, vertical='center')
for i, (_, w, _) in enumerate(cols, 1):
    s.column_dimensions[s.cell(1, i).column_letter].width = w
s.freeze_panes = 'E2'
s.auto_filter.ref = f'A1:{s.cell(1, len(cols)).column_letter}{s.max_row}'
s.row_dimensions[1].height = 32

# referências (uma linha por livro/capítulo) para tabelas dinâmicas
r2 = wb2.create_sheet('Referências')
r2.append(['ID', 'Data', 'Ano', 'Título', 'Pregador', 'Série', 'Livro', 'Testamento', 'Gênero', 'Capítulo', 'Confiança'])
for x in recs:
    for a, caps in x['refs']:
        for c in (caps or [None]):
            r2.append([x['id'], x['data'] or None, x['ano'], x['titulo'], x['pregador'], x['serie'] or None, biblia.INFO[a]['nome'],
                       biblia.INFO[a]['test'], biblia.GENERO[a], c, x['confianca']])
for c in r2[1]:
    c.font = Font(bold=True, color='FFFFFF'); c.fill = hdr
for col, w in zip('ABCDEFGHIJK', [7, 12, 7, 50, 22, 30, 18, 6, 18, 9, 13]):
    r2.column_dimensions[col].width = w
r2.auto_filter.ref = f'A1:K{r2.max_row}'
r2.freeze_panes = 'A2'

# cobertura por livro
cv = wb2.create_sheet('Cobertura por livro')
cv.append(['Livro', 'Testamento', 'Gênero', 'Capítulos no livro', 'Capítulos pregados', '% coberto', 'Nº de sermões', 'Capítulos nunca pregados'])
for a, n, t, caps, _ in biblia.LIVROS:
    feitos, ns = set(), 0
    for x in recs:
        for b, cs in x['refs']:
            if b == a:
                feitos.update(cs); ns += 1
    falta = [c for c in range(1, caps + 1) if c not in feitos]
    cv.append([n, t, biblia.GENERO[a], caps, len(feitos), round(len(feitos) / caps, 3), ns,
               biblia.fmt('', set(falta)).strip() if falta and len(falta) < caps else ('todos' if falta else '')])
    cv.cell(cv.max_row, 6).number_format = '0%'
for c in cv[1]:
    c.font = Font(bold=True, color='FFFFFF'); c.fill = hdr; c.alignment = Alignment(wrap_text=True)
for col, w in zip('ABCDEFGH', [20, 10, 18, 11, 11, 10, 11, 60]):
    cv.column_dimensions[col].width = w
cv.freeze_panes = 'A2'

info = wb2.create_sheet('Sobre')
for l in ['Sermões da Família dos Que Creem — base completa', '',
          f'{len(recs)} gravações: os 240 vídeos da lista do YouTube + os episódios do podcast (SoundCloud) que não estavam na lista.',
          '• Data: publicação no podcast (geralmente 2 a 4 dias após o culto). "Domingo provável" = domingo anterior à publicação.',
          '  Vídeos sem áudio no podcast têm só o período aproximado.',
          '• Link (abrir sermão): YouTube quando conhecido; senão o episódio do podcast.',
          '• Confiança do texto: Confirmado (anunciado no sermão ou no título) · Provável · Incerto · Temático (sem texto único) · Panorama (introdução a um livro).',
          '• Fonte do texto: Título do episódio · Transcrição do áudio (10 primeiros minutos) · NotebookLM · Descrição do episódio.',
          '• Aba "Referências": uma linha por livro/capítulo, para tabelas dinâmicas. Aba "Cobertura por livro": o que já foi e o que nunca foi pregado.']:
    info.append([l])
info.column_dimensions['A'].width = 140
info['A1'].font = Font(bold=True, size=13)
wb2.save(OUT)


# ---------- dashboard ----------
keep = ['id', 'titulo', 'pregador', 'data', 'ano', 'texto', 'confianca', 'serie', 'tipo', 'youtube', 'link', 'refs', 'livros']
d = dict(livros=[dict(ab=a, nome=n, test=t, caps=c, genero=biblia.GENERO[a]) for a, n, t, c, _ in biblia.LIVROS],
         sermoes=[{k: x.get(k) for k in keep} for x in recs])
html = open(os.path.join(AQUI, 'dashboard_template.html')).read().replace('/*DATA*/null', json.dumps(d, ensure_ascii=False, separators=(',', ':')))
open(DASH, 'w').write(html)
print(f'{len(recs)} gravações -> {os.path.relpath(OUT)} e {os.path.relpath(DASH)}')
