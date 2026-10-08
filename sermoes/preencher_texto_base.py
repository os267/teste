"""Preenche data exata e texto base na planilha sermoes_fdqc.xlsx.

Rode no SEU computador (o YouTube bloqueia servidores em nuvem):
    pip install yt-dlp openpyxl
    python preencher_texto_base.py sermoes_fdqc.xlsx

Para cada vídeo, o script:
  - pega a data de publicação exata;
  - procura referências bíblicas na descrição;
  - baixa a legenda automática em português e procura referências nos
    primeiros ~12 minutos (quando o pregador costuma anunciar o texto).
Salva em sermoes_fdqc_preenchida.xlsx. Se for interrompido, rode de novo passando
sermoes_fdqc_preenchida.xlsx: linhas que já têm data são puladas.
"""
import re
import sys
import time
import urllib.request
from collections import Counter

import yt_dlp
from openpyxl import load_workbook

LIVROS = {
    "Gênesis": r"g[êe]nesis|gn", "Êxodo": r"[êe]xodo|[êe]x", "Levítico": r"lev[íi]tico|lv",
    "Números": r"n[úu]meros|nm", "Deuteronômio": r"deuteron[ôo]mio|dt", "Josué": r"josu[ée]|js",
    "Juízes": r"ju[íi]zes|jz", "Rute": r"rute|rt", "Samuel": r"samuel|sm", "Reis": r"reis|rs",
    "Crônicas": r"cr[ôo]nicas|cr", "Esdras": r"esdras|ed", "Neemias": r"neemias|ne",
    "Ester": r"ester|et", "Jó": r"jó", "Salmos": r"salmos?|sl", "Provérbios": r"prov[ée]rbios|pv",
    "Eclesiastes": r"eclesiastes|ec", "Cantares": r"c[âa]nticos?|cantares|ct",
    "Isaías": r"isa[íi]as|is", "Jeremias": r"jeremias|jr", "Lamentações": r"lamenta[çc][õo]es|lm",
    "Ezequiel": r"ezequiel|ez", "Daniel": r"daniel|dn", "Oseias": r"os[ée]ias|os", "Joel": r"joel|jl",
    "Amós": r"am[óo]s|am", "Obadias": r"obadias|ob", "Jonas": r"jonas|jn", "Miqueias": r"miqu[ée]ias|mq",
    "Naum": r"naum|na", "Habacuque": r"habacuque|hc", "Sofonias": r"sofonias|sf", "Ageu": r"ageu|ag",
    "Zacarias": r"zacarias|zc", "Malaquias": r"malaquias|ml", "Mateus": r"mateus|mt",
    "Marcos": r"marcos|mc", "Lucas": r"lucas|lc", "João": r"jo[ãa]o|jo", "Atos": r"atos|at",
    "Romanos": r"romanos|rm", "Coríntios": r"cor[íi]ntios|co", "Gálatas": r"g[áa]latas|gl",
    "Efésios": r"ef[ée]sios|ef", "Filipenses": r"filipenses|fp", "Colossenses": r"colossenses|cl",
    "Tessalonicenses": r"tessalonicenses|ts", "Timóteo": r"tim[óo]teo|tm", "Tito": r"tito|tt",
    "Filemom": r"filemom|fm", "Hebreus": r"hebreus|hb", "Tiago": r"tiago|tg", "Pedro": r"pedro|pe",
    "Judas": r"judas|jd", "Apocalipse": r"apocalipse|ap",
}
ALT = "|".join(f"(?P<b{i}>{p})" for i, p in enumerate(LIVROS.values()))
NOMES = list(LIVROS)
# Ex.: "Mateus 5:1-12", "1 Co 12.12-27", "Jo 10"
REF_ESCRITA = re.compile(
    rf"\b(?P<n>[1-3]\s*)?(?:{ALT})\.?\s+(?P<c>\d{{1,3}})(?:\s*[:.,]\s*(?P<v>\d{{1,3}}(?:\s*[-–]\s*\d{{1,3}})?))?\b",
    re.IGNORECASE,
)
# Fala: "Mateus capítulo 5 versículo 13", "primeira carta aos Coríntios, capítulo 12"
ORD = {"primeira": "1", "primeiro": "1", "segunda": "2", "segundo": "2", "terceira": "3", "terceiro": "3"}
LIVROS_EXTENSO = "|".join(p.split("|")[0] for p in LIVROS.values())
REF_FALADA = re.compile(
    rf"(?:(?P<o>primeir[ao]|segund[ao]|terceir[ao])\s+(?:carta\s+)?(?:a|ao|aos|de|do|da)?\s*)?"
    rf"(?P<l>{LIVROS_EXTENSO})\W+(?:no\s+)?cap[íi]tulo\s+(?P<c>\d{{1,3}})"
    rf"(?:\W+(?:a\s+partir\s+do\s+)?vers[íi]culos?\s+(?P<v>\d{{1,3}}))?",
    re.IGNORECASE,
)


def nome_livro(texto):
    for nome, pat in LIVROS.items():
        if re.fullmatch(pat, texto, re.IGNORECASE):
            return nome
    return texto


def refs(texto, falada=False):
    out = []
    for m in REF_ESCRITA.finditer(texto):
        livro = next(NOMES[int(k[1:])] for k, v in m.groupdict().items() if k.startswith("b") and v)
        if len(m.group(0)) < 5 and not m.group("v"):
            continue  # evita falsos positivos tipo "at 3"
        n = (m.group("n") or "").strip()
        out.append(f"{n+' ' if n else ''}{livro} {m.group('c')}{':'+m.group('v').replace(' ', '') if m.group('v') else ''}")
    if falada:
        for m in REF_FALADA.finditer(texto):
            n = ORD.get((m.group("o") or "").lower(), "")
            out.append(f"{n+' ' if n else ''}{nome_livro(m.group('l'))} {m.group('c')}{':'+m.group('v') if m.group('v') else ''}")
    return out


def legenda_inicio(info, minutos=12):
    subs = info.get("subtitles") or {}
    auto = info.get("automatic_captions") or {}
    for fonte in (subs, auto):
        for lang in ("pt", "pt-BR", "pt-orig"):
            for f in fonte.get(lang, []):
                if f.get("ext") == "json3":
                    import json
                    data = json.load(urllib.request.urlopen(f["url"], timeout=30))
                    partes = []
                    for ev in data.get("events", []):
                        if ev.get("tStartMs", 0) > minutos * 60_000:
                            break
                        partes += [s.get("utf8", "") for s in ev.get("segs", [])]
                    return " ".join(partes)
    return ""


def main(caminho):
    saida = caminho if caminho.endswith("_preenchida.xlsx") else caminho.replace(".xlsx", "_preenchida.xlsx")
    wb = load_workbook(caminho)
    ws = wb["Sermões"]
    col = {c.value: c.column for c in ws[1]}
    for nome in ("Refs na descrição", "Refs na legenda (início)", "Descrição"):
        if nome not in col:
            col[nome] = ws.max_column + 1
            ws.cell(1, col[nome], nome)
    ydl = yt_dlp.YoutubeDL({"quiet": True, "skip_download": True, "no_warnings": True})
    for row in range(2, ws.max_row + 1):
        url = ws.cell(row, col["Link"]).value
        if not url or ws.cell(row, col["Data de publicação"]).value:
            continue
        try:
            info = ydl.extract_info(url, download=False)
        except Exception as e:
            print(row, "ERRO", e)
            continue
        d = info.get("upload_date") or ""
        ws.cell(row, col["Data de publicação"], f"{d[6:8]}/{d[4:6]}/{d[:4]}" if d else "")
        desc = info.get("description") or ""
        ws.cell(row, col["Descrição"], desc[:2000])
        r_desc = list(dict.fromkeys(refs(desc)))
        ws.cell(row, col["Refs na descrição"], "; ".join(r_desc))
        try:
            r_leg = Counter(refs(legenda_inicio(info), falada=True))
        except Exception as e:
            print(row, "legenda indisponível", e)
            r_leg = Counter()
        ws.cell(row, col["Refs na legenda (início)"], "; ".join(f"{k} ({v}x)" for k, v in r_leg.most_common(5)))
        base = ws.cell(row, col["Texto base"])
        if not base.value and r_desc:
            base.value = r_desc[0]
            ws.cell(row, col["Status do texto base"], "Confirmado (descrição)")
        elif not base.value and r_leg:
            base.value = r_leg.most_common(1)[0][0]
            ws.cell(row, col["Status do texto base"], "A confirmar")  # vindo da legenda: revisar
        print(row, d, "|", base.value or "-", "|", info.get("title"))
        if row % 10 == 0:
            wb.save(saida)
        time.sleep(2)  # evita bloqueio por excesso de requisições
    wb.save(saida)
    print("Pronto.")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "sermoes_fdqc.xlsx")
