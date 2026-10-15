"""
Genera la versión Word / Google Docs del Playbook Comercial a partir de contenido.py
y de las figuras de playbook/figuras/ (ejecute antes figuras.py).

Uso:  python3 playbook/generar_docx.py  ->  playbook/salida/Playbook_Comercial_Cremacion_Vangrupo.docx
"""
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, str(Path(__file__).parent))
import contenido as C  # noqa: E402

BASE = Path(__file__).parent
FIGURAS = BASE / "figuras"
SALIDA = BASE / "salida" / "Playbook_Comercial_Cremacion_Vangrupo.docx"

AZUL, AZUL_P, DORADO, CLARO, TEXTO, GRIS = "1B365D", "0F2537", "C5A059", "F4F7FA", "2C3E50", "7A8794"
CREMA = "FBF3E1"
TITULOS, CUERPO = "Playfair Display", "Montserrat"
ANCHO_UTIL = Cm(17)


def rgb(hex_: str) -> RGBColor:
    return RGBColor.from_string(hex_)


# ---------------------------------------------------------------------------
# utilidades de formato
# ---------------------------------------------------------------------------
def sombrear(celda, color: str) -> None:
    tcPr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def bordes(celda, **lados) -> None:
    """lados: top/left/bottom/right = (color, grosor_en_octavos_de_punto) o None para quitar."""
    tcPr = celda._tc.get_or_add_tcPr()
    tcb = OxmlElement("w:tcBorders")
    for lado in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{lado}")
        valor = lados.get(lado)
        if valor:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(valor[1]))
            el.set(qn("w:color"), valor[0])
        else:
            el.set(qn("w:val"), "nil")
        tcb.append(el)
    tcPr.append(tcb)


def margen_celda(celda, cm: float = 0.3) -> None:
    tcPr = celda._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for lado in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{lado}")
        el.set(qn("w:w"), str(int(cm * 567)))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tcPr.append(mar)


def run(par, texto: str, *, size=10.5, bold=False, italic=False, color=TEXTO, font=CUERPO):
    r = par.add_run(texto)
    r.font.name = font
    r._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = rgb(color)
    return r


def parrafo(destino, texto: str = "", *, size=10.5, bold=False, italic=False, color=TEXTO, font=CUERPO,
            alinear=None, antes=0, despues=6, interlineado=1.15):
    par = destino.add_paragraph()
    pf = par.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(antes), Pt(despues), interlineado
    if alinear:
        par.alignment = alinear
    if texto:
        run(par, texto, size=size, bold=bold, italic=italic, color=color, font=font)
    return par


def primer_parrafo(celda):
    par = celda.paragraphs[0]
    par.paragraph_format.space_after = Pt(2)
    return par


def vinetas(destino, items, *, color_marca=DORADO, size=10.5, marca="■"):
    for item in items:
        par = parrafo(destino, despues=3)
        par.paragraph_format.left_indent = Cm(0.6)
        par.paragraph_format.first_line_indent = Cm(-0.6)
        run(par, f"{marca}  ", size=size - 3, color=color_marca)
        run(par, item, size=size)


def tabla_simple(doc, filas: int, cols: int, anchos):
    t = doc.add_table(rows=filas, cols=cols)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for fila in t.rows:
        for i, celda in enumerate(fila.cells):
            celda.width = anchos[i]
            margen_celda(celda, 0.22)
    return t


def recuadro(doc, titulo: str, texto, *, fondo=CLARO, borde=DORADO, color_titulo=AZUL):
    """Caja destacada de una sola celda con borde izquierdo de color."""
    t = tabla_simple(doc, 1, 1, [ANCHO_UTIL])
    c = t.cell(0, 0)
    sombrear(c, fondo)
    bordes(c, left=(borde, 36))
    margen_celda(c, 0.35)
    par = primer_parrafo(c)
    if titulo:
        run(par, titulo.upper(), size=9, bold=True, color=color_titulo)
        par = c.add_paragraph()
        par.paragraph_format.space_after = Pt(2)
    if isinstance(texto, str):
        run(par, texto, size=10.5)
    else:
        primero = True
        for item in texto:
            p = par if primero else c.add_paragraph()
            p.paragraph_format.space_after = Pt(3)
            run(p, "■  ", size=7.5, color=borde)
            run(p, item, size=10)
            primero = False
    parrafo(doc, despues=4)
    return t


def tabla_datos(doc, encabezados, filas, anchos, *, resaltar=None):
    t = tabla_simple(doc, len(filas) + 1, len(encabezados), anchos)
    for i, h in enumerate(encabezados):
        c = t.cell(0, i)
        sombrear(c, AZUL)
        run(primer_parrafo(c), h, size=9, bold=True, color="FFFFFF")
    for r, fila in enumerate(filas, start=1):
        fondo = CREMA if resaltar and resaltar(fila) else (CLARO if r % 2 == 0 else "FFFFFF")
        for i, valor in enumerate(fila):
            c = t.cell(r, i)
            sombrear(c, fondo)
            bordes(c, bottom=("D9DEE4", 4))
            run(primer_parrafo(c), valor, size=9.5, bold=(i == 0), color=AZUL if i == 0 else TEXTO)
    parrafo(doc, despues=4)
    return t


NATIVO = False  # True: las figuras se dibujan con tablas del propio documento (sin imágenes)


def figura(doc, nombre: str, ancho=ANCHO_UTIL, pie: str = ""):
    if NATIVO:
        FIGURAS_NATIVAS[nombre](doc)
    else:
        par = parrafo(doc, alinear=WD_ALIGN_PARAGRAPH.CENTER, antes=4, despues=2)
        par.add_run().add_picture(str(FIGURAS / f"{nombre}.png"), width=ancho)
    if pie:
        parrafo(doc, pie, size=8.5, italic=True, color=GRIS, alinear=WD_ALIGN_PARAGRAPH.CENTER, despues=10)


# ---------------------------------------------------------------------------
# figuras nativas (tablas editables en Google Docs)
# ---------------------------------------------------------------------------
def _celda(c, fondo=None, margen=0.25):
    if fondo:
        sombrear(c, fondo)
    margen_celda(c, margen)
    return primer_parrafo(c)


def _nuevo_parrafo(c, alinear=None):
    par = c.add_paragraph()
    par.paragraph_format.space_after = Pt(2)
    if alinear is not None:
        par.alignment = alinear
    return par


def nat_logos(doc):
    t = tabla_simple(doc, 1, 3, [Cm(5.5)] * 3)
    for i, marca in enumerate(("La Auxiliadora", "Hogar de las Memorias", "Vangrupo")):
        c = t.cell(0, i)
        borde = (DORADO, 12)
        bordes(c, top=borde, left=borde, bottom=borde, right=borde)
        p = _celda(c, margen=0.45)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run(p, "LOGO", size=8, bold=True, color=GRIS)
        run(_nuevo_parrafo(c, WD_ALIGN_PARAGRAPH.CENTER), marca, size=12, bold=True, color=AZUL)


def nat_portafolio(doc):
    colores = [AZUL, "4F7A5A", AZUL, "4F7A5A", DORADO, "8A8F98"]
    t = tabla_simple(doc, 2, 3, [Cm(5.67)] * 3)
    for i, (p, col) in enumerate(zip(C.PRODUCTOS, colores)):
        c = t.cell(i // 3, i % 3)
        bordes(c, top=(col, 36), left=("FFFFFF", 24), right=("FFFFFF", 24), bottom=("FFFFFF", 24))
        par = _celda(c, CREMA if p.get("proximo") else CLARO, 0.3)
        if p.get("proximo"):
            run(par, "PRÓXIMO LANZAMIENTO", size=7.5, bold=True, color="B8860B")
            par = _nuevo_parrafo(c)
        run(par, p["etiqueta"].upper(), size=7.5, color=GRIS)
        run(_nuevo_parrafo(c), p["corto"], size=12.5, bold=True, color=AZUL_P)
        run(_nuevo_parrafo(c), p["precio"], size=16, bold=True, color=col)
        run(_nuevo_parrafo(c), p["pago"], size=8.5)


def nat_embudo(doc):
    n = len(C.EMBUDO)
    t = tabla_simple(doc, n, 3, [Cm(1.1), Cm(4.6), Cm(11.3)])
    tonos = ["1B365D", "2E4A70", "425E83", "567296", "6A86A9", "7E9ABC", "92AECF"]
    for i, (etapa, desc) in enumerate(C.EMBUDO):
        c0, c1, c2 = t.cell(i, 0), t.cell(i, 1), t.cell(i, 2)
        for c in (c0, c1, c2):
            bordes(c, bottom=("FFFFFF", 16))
        p0 = _celda(c0, DORADO)
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run(p0, str(i + 1), size=12, bold=True, color=AZUL_P)
        run(_celda(c1, tonos[i]), etapa, size=11.5, bold=True, color="FFFFFF")
        run(_celda(c2, CLARO if i % 2 == 0 else "FFFFFF"), desc, size=10)


def nat_conversacion(doc):
    t = tabla_simple(doc, 2, len(C.CONVERSACION), [Cm(3.4)] * len(C.CONVERSACION))
    for i, (paso, desc) in enumerate(C.CONVERSACION):
        c0, c1 = t.cell(0, i), t.cell(1, i)
        bordes(c0, right=("FFFFFF", 24), bottom=(DORADO, 24))
        p = _celda(c0, AZUL)
        run(p, f"{i + 1}  ", size=13, bold=True, color=DORADO)
        run(p, paso, size=10.5, bold=True, color="FFFFFF")
        run(_celda(c1), desc, size=9)


def nat_equipos(doc):
    t = tabla_simple(doc, 3, 3, [Cm(5.67)] * 3)
    for i, (eq, desc, prod) in enumerate(C.EQUIPOS):
        cab, mid, pie = t.cell(0, i), t.cell(1, i), t.cell(2, i)
        bordes(cab, bottom=(DORADO, 24), left=("FFFFFF", 24), right=("FFFFFF", 24))
        run(_celda(cab, AZUL), eq, size=10.5, bold=True, color="FFFFFF")
        bordes(mid, left=("FFFFFF", 24), right=("FFFFFF", 24))
        run(_celda(mid, CLARO), desc, size=9.5)
        borde = (DORADO, 12)
        bordes(pie, top=borde, left=borde, bottom=borde, right=borde)
        p = _celda(pie)
        run(p, "▼ PRODUCTO FOCO  ", size=7.5, bold=True, color=DORADO)
        run(_nuevo_parrafo(pie), prod, size=9.5, bold=True, color=AZUL)


def nat_ciclo(doc):
    pasos = [
        ("Inhumación", "Urna biodegradable en el jardín, con placa de identificación."),
        ("3 años", "La familia visita un lugar señalado; la urna se integra a la tierra."),
        ("Placa memorial", "El nombre pasa al monumento del jardín (incluida en el Ecológico y en Mascotas US$ 550)."),
        ("Ciclo natural", "El espacio se reutiliza: el jardín sigue vivo para otras familias."),
    ]
    t = tabla_simple(doc, 2, 4, [Cm(4.25)] * 4)
    for i, (tit, desc) in enumerate(pasos):
        c0, c1 = t.cell(0, i), t.cell(1, i)
        bordes(c0, right=("FFFFFF", 24))
        p = _celda(c0, "4F7A5A")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run(p, f"{i + 1}", size=16, bold=True, color="FFFFFF")
        run(_nuevo_parrafo(c0, WD_ALIGN_PARAGRAPH.CENTER), tit, size=10.5, bold=True, color="FFFFFF")
        p1 = _celda(c1, CLARO)
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run(p1, desc, size=9)


def nat_objeciones(doc):
    t = tabla_simple(doc, 1, 3, [Cm(5.67)] * 3)
    for i, (paso, desc) in enumerate(C.METODO_OBJECIONES):
        c = t.cell(0, i)
        bordes(c, top=(DORADO, 36), left=("FFFFFF", 24), right=("FFFFFF", 24))
        p = _celda(c, CLARO, 0.35)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run(p, paso[0], size=30, bold=True, color=AZUL, font=TITULOS)
        run(_nuevo_parrafo(c, WD_ALIGN_PARAGRAPH.CENTER), f"{i + 1}. {paso}", size=12, bold=True, color=AZUL_P)
        run(_nuevo_parrafo(c, WD_ALIGN_PARAGRAPH.CENTER), desc, size=9.5)
    parrafo(doc, despues=2)
    recuadro(doc, "", C.PRINCIPIO_OBJECIONES)


def nat_comisiones(doc):
    maximo = max(v for _, v in C.COMISIONES_VALORES)
    t = tabla_simple(doc, len(C.COMISIONES_VALORES), 2, [Cm(4.5), Cm(12.5)])
    for i, (nombre, valor) in enumerate(C.COMISIONES_VALORES):
        c0, c1 = t.cell(i, 0), t.cell(i, 1)
        p0 = _celda(c0, margen=0.15)
        p0.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run(p0, nombre, size=10, bold=True, color=AZUL_P)
        p1 = _celda(c1, margen=0.15)
        run(p1, "█" * max(1, round(valor / maximo * 26)), size=10, color=AZUL)
        monto = f"{valor:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
        run(p1, f"  US$ {monto}", size=10, bold=True, color=AZUL)


FIGURAS_NATIVAS = {
    "00_logos": nat_logos, "01_portafolio": nat_portafolio, "02_embudo": nat_embudo,
    "03_conversacion": nat_conversacion, "04_objeciones": nat_objeciones, "05_comisiones": nat_comisiones,
    "06_ciclo_ecologico": nat_ciclo, "07_equipos": nat_equipos,
}


def salto(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def seccion(doc, numero: str, titulo: str, intro: str = ""):
    """Encabezado de sección: estilo Heading 1 (para el esquema de Google Docs) con número dorado."""
    h = doc.add_paragraph(style="Heading 1")
    h.paragraph_format.space_before, h.paragraph_format.space_after = Pt(0), Pt(4)
    run(h, f"{numero}  ", size=26, bold=True, color=DORADO, font=TITULOS)
    run(h, titulo, size=24, bold=True, color=AZUL, font=TITULOS)
    linea = parrafo(doc, despues=10)
    pPr = linea._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", "12"), ("color", DORADO), ("space", "1")):
        bottom.set(qn(f"w:{k}"), v)
    pbdr.append(bottom)
    pPr.append(pbdr)
    if intro:
        parrafo(doc, intro, size=11, italic=True, color=GRIS, despues=10)


def subtitulo(doc, texto: str, color=AZUL, size=14, antes=10):
    h = doc.add_paragraph(style="Heading 2")
    h.paragraph_format.space_before, h.paragraph_format.space_after = Pt(antes), Pt(4)
    run(h, texto, size=size, bold=True, color=color, font=TITULOS)
    return h


def etiqueta(doc, texto: str):
    parrafo(doc, texto.upper(), size=8.5, bold=True, color=DORADO, antes=6, despues=2)


# ---------------------------------------------------------------------------
# documento
# ---------------------------------------------------------------------------
def estilos(doc):
    st = doc.styles["Normal"]
    st.font.name, st.font.size = CUERPO, Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), CUERPO)
    for nombre in ("Heading 1", "Heading 2"):
        s = doc.styles[nombre]
        s.font.name = TITULOS
        s.font.color.rgb = rgb(AZUL)
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(27.94), Cm(21.59)  # Carta
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin, sec.bottom_margin = Cm(2), Cm(2)


def portada(doc):
    figura(doc, "00_logos", ancho=Cm(15))
    parrafo(doc, despues=40)
    t = tabla_simple(doc, 1, 1, [ANCHO_UTIL])
    c = t.cell(0, 0)
    sombrear(c, AZUL)
    bordes(c, bottom=(DORADO, 48))
    margen_celda(c, 0.9)
    p = primer_parrafo(c)
    run(p, C.CAMPANA.upper(), size=10, bold=True, color=DORADO)
    p = c.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    run(p, C.TITULO, size=40, bold=True, color="FFFFFF", font=TITULOS)
    p = c.add_paragraph()
    run(p, C.SUBTITULO, size=16, color="E8EDF3", font=TITULOS)
    p = c.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    run(p, C.MARCAS, size=10.5, bold=True, color=DORADO)
    parrafo(doc, despues=30)
    parrafo(doc, f"«{C.LEMA}»", size=15, italic=True, color=AZUL, font=TITULOS,
            alinear=WD_ALIGN_PARAGRAPH.CENTER, despues=60)
    parrafo(doc, C.VERSION, size=9, color=GRIS, alinear=WD_ALIGN_PARAGRAPH.CENTER)
    salto(doc)


def como_usar(doc):
    subtitulo(doc, "Cómo usar este playbook", size=18, antes=0)
    parrafo(doc, C.COMO_USAR, despues=10)
    recuadro(doc, "Pendientes de validación (versión 1.0)", C.PENDIENTES, fondo=CREMA, borde=DORADO)
    subtitulo(doc, "Contenido", size=18)
    t = tabla_simple(doc, len(C.INDICE), 2, [Cm(2), Cm(15)])
    for i, (num, tit) in enumerate(C.INDICE):
        c0, c1 = t.cell(i, 0), t.cell(i, 1)
        for c in (c0, c1):
            bordes(c, bottom=("D9DEE4", 4))
        run(primer_parrafo(c0), num, size=18, bold=True, color=DORADO, font=TITULOS)
        run(primer_parrafo(c1), tit, size=12, bold=True, color=AZUL)
    salto(doc)


def s01(doc):
    seccion(doc, "01", "Quiénes somos y lo que vendemos")
    parrafo(doc, C.QUIENES_SOMOS, despues=10)
    figura(doc, "01_portafolio", pie="Figura 1 · El portafolio en un vistazo. Precios sugeridos, IVA incluido.")
    subtitulo(doc, "¿Por qué ahora?")
    t = tabla_simple(doc, 2, 2, [Cm(8.5), Cm(8.5)])
    for i, (tit, txt) in enumerate(C.POR_QUE_AHORA):
        c = t.cell(i // 2, i % 2)
        sombrear(c, CLARO)
        bordes(c, top=(DORADO, 24), left=("FFFFFF", 24), right=("FFFFFF", 24), bottom=("FFFFFF", 24))
        margen_celda(c, 0.35)
        run(primer_parrafo(c), tit, size=11.5, bold=True, color=AZUL)
        run(c.add_paragraph(), txt, size=10)
    parrafo(doc, despues=4)
    subtitulo(doc, "Dos momentos de venta")
    for tit, txt in C.MODOS_DE_VENTA:
        p = parrafo(doc, despues=5)
        run(p, f"{tit}. ", bold=True, color=AZUL)
        run(p, txt)
    recuadro(doc, "Regla de oro", C.REGLA_DE_ORO, fondo=AZUL, borde=DORADO, color_titulo=DORADO)
    # el texto del recuadro azul debe ir en blanco
    salto(doc)


def s02(doc):
    seccion(doc, "02", "A quién le vendemos",
            "Cinco perfiles de cliente. Identifique el perfil en los primeros minutos y adapte el tono.")
    tabla_datos(doc, ["Perfil", "Quién es", "Qué busca", "Tono", "Producto ideal"],
                [list(p) for p in C.PERFILES], [Cm(3.1), Cm(3.3), Cm(3.8), Cm(3.1), Cm(3.7)])
    recuadro(doc, "Filtro de calificación: N · O · A", C.FILTRO_CALIFICACION)
    salto(doc)


def s03(doc):
    seccion(doc, "03", "El proceso de venta",
            "Un camino claro, del primer contacto a la posventa. Cada etapa tiene un objetivo.")
    figura(doc, "02_embudo", ancho=Cm(16), pie="Figura 2 · Las 7 etapas del embudo de prenecesidad.")
    figura(doc, "03_conversacion", pie="Figura 3 · Toda conversación sigue cinco pasos.")
    subtitulo(doc, "Qué producto lleva cada equipo")
    figura(doc, "07_equipos", pie="Figura 4 · Enfoque de producto por canal.")
    subtitulo(doc, "Rituales del equipo")
    tabla_datos(doc, ["Frecuencia", "Ritual"], [list(r) for r in C.RITUALES], [Cm(3.5), Cm(13.5)])
    salto(doc)


def ficha_producto(doc, p):
    proximo = p.get("proximo")
    # banda de título
    t = tabla_simple(doc, 1, 2, [Cm(11.5), Cm(5.5)])
    izq, der = t.cell(0, 0), t.cell(0, 1)
    color = "8A8F98" if proximo else AZUL
    for c in (izq, der):
        sombrear(c, color)
        bordes(c, bottom=(DORADO, 36))
        margen_celda(c, 0.4)
    run(primer_parrafo(izq), p["etiqueta"].upper(), size=8.5, bold=True, color=DORADO)
    run(izq.add_paragraph(), p["nombre"], size=18, bold=True, color="FFFFFF", font=TITULOS)
    pd = primer_parrafo(der)
    pd.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run(pd, "PRÓXIMO LANZAMIENTO" if proximo else "PRECIO SUGERIDO", size=8.5, bold=True, color=DORADO)
    pd2 = der.add_paragraph()
    pd2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run(pd2, p["precio"], size=18, bold=True, color="FFFFFF")
    parrafo(doc, p["pago"], size=9.5, italic=True, color=GRIS, antes=4, despues=8)

    p_que = parrafo(doc, despues=6)
    run(p_que, "Qué es. ", bold=True, color=AZUL)
    run(p_que, p["que_es"])
    p_quien = parrafo(doc, despues=8)
    run(p_quien, "Para quién. ", bold=True, color=AZUL)
    run(p_quien, p["para_quien"])

    t = tabla_simple(doc, 1, 2, [Cm(8.5), Cm(8.5)])
    for c, tit, items, marca in ((t.cell(0, 0), "Qué incluye", p["incluye"], "✓"),
                                 (t.cell(0, 1), "Reglas del producto", p["reglas"], "■")):
        sombrear(c, CLARO)
        bordes(c, top=(DORADO, 24), left=("FFFFFF", 24), right=("FFFFFF", 24))
        margen_celda(c, 0.3)
        run(primer_parrafo(c), tit.upper(), size=8.5, bold=True, color=AZUL)
        for item in items:
            par = c.add_paragraph()
            par.paragraph_format.space_after = Pt(2)
            par.paragraph_format.left_indent = Cm(0.45)
            par.paragraph_format.first_line_indent = Cm(-0.45)
            run(par, f"{marca}  ", size=9, bold=True, color="4F7A5A" if marca == "✓" else DORADO)
            run(par, item, size=9.5)
    parrafo(doc, despues=4)

    etiqueta(doc, "Speech sugerido")
    recuadro(doc, "", p["speech"], fondo=CREMA, borde=DORADO)
    etiqueta(doc, "Preguntas de descubrimiento")
    vinetas(doc, p["preguntas"], marca="?", size=10)
    etiqueta(doc, "Frase de cierre")
    parrafo(doc, p["cierre"], italic=True, color=AZUL, despues=6)
    etiqueta(doc, "Objeción típica")
    parrafo(doc, p["objecion"], size=10, despues=6)
    recuadro(doc, "Tip del asesor", p["tip"], fondo=CLARO, borde="4F7A5A", color_titulo="4F7A5A")


def s04(doc):
    seccion(doc, "04", "Producto por producto: speech y reglas",
            "Para cada producto: qué es, qué incluye, sus reglas y cómo presentarlo.")
    for i, p in enumerate(C.PRODUCTOS):
        if i:
            salto(doc)
        ficha_producto(doc, p)
        if p["corto"] == "Jardín Ecológico":
            figura(doc, "06_ciclo_ecologico", pie="Figura 5 · El ciclo de 3 años del Jardín Ecológico "
                                                  "y de Mascotas, explicado al cliente.")
    salto(doc)


def s05(doc):
    seccion(doc, "05", "Manejo de objeciones",
            "Una objeción no es un «no»: es una duda que aún no resolvimos.")
    figura(doc, "04_objeciones", pie="Figura 6 · Método V·R·A: Validar, Reposicionar, Avanzar.")
    for i, (obj, val, rep, ava) in enumerate(C.OBJECIONES, start=1):
        t = tabla_simple(doc, 4, 2, [Cm(3.2), Cm(13.8)])
        cab = t.cell(0, 0).merge(t.cell(0, 1))
        sombrear(cab, AZUL)
        run(primer_parrafo(cab), f"{i}.  {obj}", size=11, bold=True, color="FFFFFF")
        for fila, (letra, nombre, texto) in enumerate(
                (("V", "Validar", val), ("R", "Reposicionar", rep), ("A", "Avanzar", ava)), start=1):
            c0, c1 = t.cell(fila, 0), t.cell(fila, 1)
            sombrear(c0, CLARO)
            bordes(c0, bottom=("D9DEE4", 4))
            bordes(c1, bottom=("D9DEE4", 4))
            p0 = primer_parrafo(c0)
            run(p0, f"{letra}  ", size=12, bold=True, color=DORADO, font=TITULOS)
            run(p0, nombre, size=9, bold=True, color=AZUL)
            texto = texto if letra == "R" else f"«{texto}»"
            run(primer_parrafo(c1), texto, size=10, italic=(letra != "R"))
        parrafo(doc, despues=6)
    salto(doc)


def s06(doc):
    seccion(doc, "06", "Precios, reglas y condiciones",
            "Fuente oficial: Ficha Técnica de Cremación vigente. Precios en US$ con IVA incluido.")
    tabla_datos(doc, ["Producto", "Precio", "Forma de pago", "Mantenimiento", "Plazo / condición"],
                [list(f) for f in C.TABLA_PRECIOS], [Cm(3.4), Cm(2.6), Cm(4), Cm(2.8), Cm(4.2)],
                resaltar=lambda f: f[0] == "Jardín Ceremonial")
    subtitulo(doc, "Reglas comerciales")
    tabla_datos(doc, ["Tema", "Regla"], [list(r) for r in C.REGLAS_COMERCIALES], [Cm(4), Cm(13)])
    recuadro(doc, "El descuento no es la salida", C.DESCUENTO_NO_ES_SALIDA, fondo=CREMA)
    salto(doc)


def s07(doc):
    seccion(doc, "07", "Política de comisiones")
    recuadro(doc, "Comisión única de 8%",
             "Todos los productos comisionan 8% sobre el precio sin IVA. Confirmar con la Gerencia "
             "Comercial antes de aplicar.",
             fondo=CREMA, borde="B8860B", color_titulo="B8860B")
    figura(doc, "05_comisiones", pie="Figura 7 · Comisión por venta (8% sobre el precio sin IVA).")
    tabla_datos(doc, ["Producto", "%", "Comisión (sin IVA)", "Cuándo se paga", "Estorno"],
                [list(f) for f in C.COMISIONES], [Cm(3.4), Cm(1.4), Cm(2.6), Cm(5.6), Cm(4)],
                resaltar=lambda f: f[0] == "Jardín Ceremonial")
    subtitulo(doc, "Reglas de comisión")
    vinetas(doc, C.REGLAS_COMISION)
    parrafo(doc, despues=4)
    recuadro(doc, "Por qué la comisión está ligada al cobro", C.COMISION_CUIDADO)
    salto(doc)


def s08(doc):
    seccion(doc, "08", "Herramientas de apoyo")
    tabla_datos(doc, ["Herramienta", "Para qué sirve"], [list(h) for h in C.HERRAMIENTAS], [Cm(5), Cm(12)])
    subtitulo(doc, "Checklist antes de cerrar")
    for item in C.CHECKLIST_CIERRE:
        p = parrafo(doc, despues=4)
        run(p, "☐  ", size=12, color=AZUL)
        run(p, item)
    parrafo(doc, despues=24)
    t = tabla_simple(doc, 1, 1, [ANCHO_UTIL])
    c = t.cell(0, 0)
    sombrear(c, AZUL)
    bordes(c, bottom=(DORADO, 48))
    margen_celda(c, 0.8)
    p = primer_parrafo(c)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(p, C.CIERRE_FRASE, size=14, italic=True, color="FFFFFF", font=TITULOS)
    p = c.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    run(p, C.MARCAS, size=9.5, bold=True, color=DORADO)


def pie_de_pagina(doc):
    for sec in doc.sections:
        sec.different_first_page_header_footer = True
        p = sec.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run(p, f"{C.TITULO} · {C.MARCAS} · Uso interno", size=8, color=GRIS)


def blanquear_recuadro_azul(doc):
    """Los recuadros con fondo azul llevan el texto del cuerpo en blanco."""
    for t in doc.tables:
        c = t.cell(0, 0)
        shd = c._tc.tcPr.find(qn("w:shd")) if c._tc.tcPr is not None else None
        if shd is not None and shd.get(qn("w:fill")) == AZUL and len(t.rows) == 1 and len(t.columns) == 1:
            for par in c.paragraphs:
                for r in par.runs:
                    if r.font.color.rgb == rgb(TEXTO):
                        r.font.color.rgb = rgb("FFFFFF")


def main():
    global NATIVO, SALIDA
    if "--nativo" in sys.argv:  # versión liviana para convertir en Google Docs
        NATIVO = True
        SALIDA = SALIDA.with_name(SALIDA.stem + "_GoogleDocs.docx")
    doc = Document()
    estilos(doc)
    portada(doc)
    como_usar(doc)
    for s in (s01, s02, s03, s04, s05, s06, s07, s08):
        s(doc)
    pie_de_pagina(doc)
    blanquear_recuadro_azul(doc)
    SALIDA.parent.mkdir(exist_ok=True)
    doc.save(SALIDA)
    if NATIVO:
        quitar_partes_opcionales(SALIDA)
    print(f"✔ {SALIDA} ({SALIDA.stat().st_size:,} bytes)")


def quitar_partes_opcionales(ruta: Path) -> None:
    """Elimina stylesWithEffects.xml (opcional) para que el archivo pese menos."""
    import re
    import zipfile
    datos = {}
    with zipfile.ZipFile(ruta) as z:
        for info in z.infolist():
            datos[info.filename] = z.read(info.filename)
    datos.pop("word/stylesWithEffects.xml", None)
    datos["[Content_Types].xml"] = re.sub(rb'<Override PartName="/word/stylesWithEffects.xml"[^>]*/>', b"",
                                          datos["[Content_Types].xml"])
    datos["word/_rels/document.xml.rels"] = re.sub(rb'<Relationship [^>]*stylesWithEffects[^>]*/>', b"",
                                                   datos["word/_rels/document.xml.rels"])
    with zipfile.ZipFile(ruta, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as z:
        for nombre, contenido in datos.items():
            z.writestr(nombre, contenido)


if __name__ == "__main__":
    main()
