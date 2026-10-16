"""
Genera la versión presentación (16:9) del Playbook Comercial para importar en Canva,
a partir de contenido.py y de las figuras de playbook/figuras/ (ejecute antes figuras.py).

Uso:  python3 playbook/generar_pptx.py  ->  playbook/salida/Playbook_Comercial_Cremacion_Vangrupo.pptx
"""
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

sys.path.insert(0, str(Path(__file__).parent))
import contenido as C  # noqa: E402

BASE = Path(__file__).parent
FIGURAS = BASE / "figuras"
SALIDA = BASE / "salida" / "Playbook_Comercial_Cremacion_Vangrupo.pptx"

import paleta as P  # noqa: E402

AZUL, AZUL_P, DORADO, CLARO, TEXTO, GRIS = P.PRINCIPAL, P.PRINCIPAL_O, P.ACENTO, P.CLARO, P.TEXTO, P.GRIS
VERDE, CREMA, BLANCO = P.VERDE, P.CREMA, "FFFFFF"
TITULOS, CUERPO = "Playfair Display", "Montserrat"
W, H = Inches(13.333), Inches(7.5)
MX = Inches(0.6)  # margen lateral


def rgb(h):
    return RGBColor.from_string(h)


# ---------------------------------------------------------------------------
# primitivas
# ---------------------------------------------------------------------------
def caja(s, x, y, w, h, fondo=None, borde=None, forma=MSO_SHAPE.RECTANGLE, grosor=1.0):
    sh = s.shapes.add_shape(forma, x, y, w, h)
    if fondo:
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(fondo)
    else:
        sh.fill.background()
    if borde:
        sh.line.color.rgb = rgb(borde)
        sh.line.width = Pt(grosor)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    if forma == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = 0.08
    return sh


def texto(s, x, y, w, h, contenido, *, size=14, color=TEXTO, bold=False, italic=False, font=CUERPO,
          alinear=PP_ALIGN.LEFT, ancla=MSO_ANCHOR.TOP, interlineado=1.15, espacio=4):
    """contenido: str o lista de párrafos; cada párrafo es str o lista de (texto, {estilo})."""
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = ancla
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    parrafos = contenido if isinstance(contenido, list) else [contenido]
    for i, par in enumerate(parrafos):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = alinear
        p.line_spacing = interlineado
        p.space_after = Pt(espacio)
        trozos = par if isinstance(par, list) else [(par, {})]
        for t, est in trozos:
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = est.get("font", font)
            f.size = Pt(est.get("size", size))
            f.bold = est.get("bold", bold)
            f.italic = est.get("italic", italic)
            f.color.rgb = rgb(est.get("color", color))
    return tb


def imagen(s, nombre, x, y, w=None, h=None):
    ruta = FIGURAS / f"{nombre}.png"
    ancho, alto = Image.open(ruta).size
    if w and h:  # encajar dentro de la caja, centrado
        esc = min(w / ancho, h / alto)
        ww, hh = int(ancho * esc), int(alto * esc)
        return s.shapes.add_picture(str(ruta), x + (w - ww) // 2, y + (h - hh) // 2, ww, hh)
    return s.shapes.add_picture(str(ruta), x, y, w, h)


def nueva(prs, fondo=BLANCO):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(fondo)
    return s


def pie(s, num, oscuro=False):
    color = P.TINTE if oscuro else GRIS
    texto(s, MX, H - Inches(0.45), W - 2 * MX - Inches(1), Inches(0.3), f"{C.TITULO}  ·  {C.MARCAS}",
          size=9, color=color)
    texto(s, W - MX - Inches(1), H - Inches(0.45), Inches(1), Inches(0.3), str(num),
          size=9, color=color, alinear=PP_ALIGN.RIGHT)


def encabezado(s, seccion, titulo, intro=""):
    caja(s, 0, 0, Inches(0.18), H, fondo=AZUL)
    if seccion:
        texto(s, MX, Inches(0.4), Inches(8), Inches(0.35), seccion.upper(), size=11, bold=True, color=DORADO)
    texto(s, MX, Inches(0.68), W - 2 * MX, Inches(0.75), titulo, size=30, bold=True, color=AZUL, font=TITULOS)
    caja(s, MX, Inches(1.45), Inches(1.1), Inches(0.05), fondo=DORADO)
    if intro:
        texto(s, MX, Inches(1.58), W - 2 * MX, Inches(0.45), intro, size=13, italic=True, color=GRIS)


def tarjeta(s, x, y, w, h, titulo, cuerpo, *, fondo=CLARO, acento=DORADO, size_t=14, size_c=12,
            color_t=AZUL, color_c=TEXTO):
    caja(s, x, y, w, h, fondo=fondo)
    caja(s, x, y, w, Inches(0.06), fondo=acento)
    texto(s, x + Inches(0.2), y + Inches(0.18), w - Inches(0.4), Inches(0.62), titulo,
          size=size_t, bold=True, color=color_t, ancla=MSO_ANCHOR.MIDDLE, interlineado=1.0)
    texto(s, x + Inches(0.2), y + Inches(0.82), w - Inches(0.4), h - Inches(0.92), cuerpo,
          size=size_c, color=color_c)


def lista(items, marca="■", color_marca=DORADO, size=13):
    return [[(f"{marca}  ", {"color": color_marca, "size": size - 3, "bold": True}), (it, {"size": size})]
            for it in items]


def tabla(s, x, y, w, encabezados, filas, anchos, *, size=11, alto_fila=Inches(0.42), resaltar=None):
    t = s.shapes.add_table(len(filas) + 1, len(encabezados), x, y, w,
                           alto_fila * (len(filas) + 1)).table
    total = sum(anchos)
    for i, a in enumerate(anchos):
        t.columns[i].width = Emu(int(w * a / total))
    for r in range(len(filas) + 1):
        t.rows[r].height = alto_fila
        for c in range(len(encabezados)):
            celda = t.cell(r, c)
            celda.margin_left = celda.margin_right = Inches(0.08)
            celda.margin_top = celda.margin_bottom = Inches(0.04)
            celda.vertical_anchor = MSO_ANCHOR.MIDDLE
            if r == 0:
                valor, fondo, color, negrita = encabezados[c], AZUL, BLANCO, True
            else:
                fila = filas[r - 1]
                valor = fila[c]
                fondo = CREMA if resaltar and resaltar(fila) else (CLARO if r % 2 == 0 else BLANCO)
                color, negrita = (AZUL, True) if c == 0 else (TEXTO, False)
            celda.fill.solid()
            celda.fill.fore_color.rgb = rgb(fondo)
            tf = celda.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = ""
            run = p.add_run()
            run.text = valor
            run.font.name, run.font.size, run.font.bold = CUERPO, Pt(size), negrita
            run.font.color.rgb = rgb(color)
    return t


# ---------------------------------------------------------------------------
# diapositivas
# ---------------------------------------------------------------------------
def portada(prs):
    s = nueva(prs, AZUL)
    caja(s, 0, H - Inches(0.25), W, Inches(0.25), fondo=DORADO)
    texto(s, MX + Inches(0.2), Inches(1.0), Inches(9), Inches(0.4), C.CAMPANA.upper(),
          size=13, bold=True, color=DORADO)
    texto(s, MX + Inches(0.2), Inches(1.5), Inches(11), Inches(1.3), C.TITULO,
          size=60, bold=True, color=BLANCO, font=TITULOS)
    texto(s, MX + Inches(0.2), Inches(2.8), Inches(11), Inches(0.7), C.SUBTITULO,
          size=26, color="F1ECF7", font=TITULOS)
    caja(s, MX + Inches(0.25), Inches(3.65), Inches(1.4), Inches(0.06), fondo=DORADO)
    texto(s, MX + Inches(0.2), Inches(3.9), Inches(11), Inches(0.5), f"«{C.LEMA}»",
          size=16, italic=True, color=BLANCO, font=TITULOS)
    # espacios para logos (reemplazar en Canva)
    ancho, alto, gap = Inches(3.3), Inches(1.15), Inches(0.35)
    x0 = MX + Inches(0.2)
    for i, marca in enumerate(("La Auxiliadora", "Hogar de las Memorias", "Vangrupo")):
        x = x0 + i * (ancho + gap)
        b = caja(s, x, Inches(5.1), ancho, alto, fondo=BLANCO, forma=MSO_SHAPE.ROUNDED_RECTANGLE)
        b.name = f"LOGO {marca}"
        texto(s, x, Inches(5.1), ancho, alto,
              [[("LOGO", {"size": 9, "bold": True, "color": GRIS})], [(marca, {"size": 15, "bold": True})]],
              color=AZUL, alinear=PP_ALIGN.CENTER, ancla=MSO_ANCHOR.MIDDLE, espacio=0)
    texto(s, MX + Inches(0.2), Inches(6.55), Inches(11), Inches(0.35), C.VERSION, size=10, color=P.TINTE)


def como_usar(prs, n):
    s = nueva(prs)
    encabezado(s, "Antes de empezar", "Cómo usar este playbook")
    texto(s, MX, Inches(1.75), Inches(6.2), Inches(2.0), C.COMO_USAR, size=14)
    caja(s, MX, Inches(3.75), Inches(6.2), Inches(2.9), fondo=CREMA)
    caja(s, MX, Inches(3.75), Inches(0.08), Inches(2.9), fondo=DORADO)
    texto(s, MX + Inches(0.3), Inches(3.9), Inches(5.7), Inches(0.4), "PENDIENTES DE VALIDACIÓN",
          size=11, bold=True, color=AZUL)
    texto(s, MX + Inches(0.3), Inches(4.3), Inches(5.7), Inches(2.3), lista(C.PENDIENTES, size=12))
    x = Inches(7.35)
    texto(s, x, Inches(1.75), Inches(5.3), Inches(0.5), "Contenido", size=20, bold=True, color=AZUL, font=TITULOS)
    for i, (num, tit) in enumerate(C.INDICE):
        y = Inches(2.35) + i * Inches(0.53)
        texto(s, x, y, Inches(0.8), Inches(0.5), num, size=22, bold=True, color=DORADO, font=TITULOS)
        texto(s, x + Inches(0.8), y + Inches(0.08), Inches(4.6), Inches(0.45), tit, size=14, bold=True, color=AZUL)
        caja(s, x, y + Inches(0.5), Inches(5.3), Emu(9525), fondo=P.LINEA)
    pie(s, n)


def separador(prs, n, num, titulo, frase):
    s = nueva(prs, AZUL)
    texto(s, MX + Inches(0.4), Inches(1.55), Inches(3.5), Inches(2.1), num, size=110, bold=True,
          color=DORADO, font=TITULOS)
    texto(s, MX + Inches(0.4), Inches(3.75), Inches(11.5), Inches(1.0), titulo, size=40, bold=True,
          color=BLANCO, font=TITULOS)
    caja(s, MX + Inches(0.45), Inches(4.85), Inches(1.4), Inches(0.06), fondo=DORADO)
    texto(s, MX + Inches(0.4), Inches(5.05), Inches(11), Inches(0.8), frase, size=16, italic=True, color=P.TINTE)
    pie(s, n, oscuro=True)


def quienes(prs, n):
    s = nueva(prs)
    encabezado(s, "01 · Quiénes somos", "No vendemos parcelas: acompañamos familias")
    texto(s, MX, Inches(1.75), Inches(12), Inches(1.4), C.QUIENES_SOMOS, size=15)
    w, gap = Inches(2.88), Inches(0.2)
    for i, (tit, txt) in enumerate(C.POR_QUE_AHORA):
        tarjeta(s, MX + i * (w + gap), Inches(3.45), w, Inches(1.9), tit, txt, size_c=12)
    caja(s, MX, Inches(5.65), W - 2 * MX, Inches(1.05), fondo=AZUL)
    texto(s, MX + Inches(0.3), Inches(5.65), W - 2 * MX - Inches(0.6), Inches(1.05),
          [[("REGLA DE ORO   ", {"bold": True, "color": DORADO, "size": 12}),
            (C.REGLA_DE_ORO, {"color": BLANCO, "size": 14, "italic": True})]], ancla=MSO_ANCHOR.MIDDLE)
    pie(s, n)


def portafolio(prs, n):
    s = nueva(prs)
    encabezado(s, "01 · Lo que vendemos", "El portafolio en un vistazo")
    imagen(s, "01_portafolio", MX, Inches(1.65), W - 2 * MX, Inches(5.2))
    pie(s, n)


def momentos(prs, n):
    s = nueva(prs)
    encabezado(s, "01 · Lo que vendemos", "Dos momentos de venta, dos actitudes")
    w = Inches(5.95)
    for i, (tit, txt) in enumerate(C.MODOS_DE_VENTA):
        tarjeta(s, MX + i * (w + Inches(0.25)), Inches(1.9), w, Inches(2.4), tit, txt,
                acento=AZUL if i == 0 else VERDE, size_t=18, size_c=15)
    texto(s, MX, Inches(4.7), W - 2 * MX, Inches(0.5), "Qué producto lleva cada equipo",
          size=18, bold=True, color=AZUL, font=TITULOS)
    w3 = Inches(3.9)
    for i, (eq, desc, prod) in enumerate(C.EQUIPOS):
        x = MX + i * (w3 + Inches(0.2))
        caja(s, x, Inches(5.25), w3, Inches(1.55), fondo=CLARO)
        texto(s, x + Inches(0.15), Inches(5.32), w3 - Inches(0.3), Inches(1.45),
              [[(eq, {"bold": True, "color": AZUL, "size": 12})], [(prod, {"size": 11})]], espacio=3)
    pie(s, n)


def perfiles(prs, n):
    s = nueva(prs)
    encabezado(s, "02 · A quién le vendemos", "Cinco perfiles de cliente",
               "Identifique el perfil en los primeros minutos y adapte el tono.")
    w, gap = Inches(2.28), Inches(0.17)
    for i, (nom, quien, busca, tono, prod) in enumerate(C.PERFILES):
        x, y = MX + i * (w + gap), Inches(2.2)
        caja(s, x, y, w, Inches(3.55), fondo=CLARO)
        caja(s, x, y, w, Inches(0.75), fondo=AZUL)
        texto(s, x + Inches(0.12), y, w - Inches(0.24), Inches(0.75), nom, size=14, bold=True,
              color=BLANCO, ancla=MSO_ANCHOR.MIDDLE)
        texto(s, x + Inches(0.12), y + Inches(0.85), w - Inches(0.24), Inches(2.65), [
            [(quien, {"italic": True, "color": GRIS, "size": 10.5})],
            [("Busca: ", {"bold": True, "color": AZUL, "size": 10.5}), (busca, {"size": 10.5})],
            [("Tono: ", {"bold": True, "color": AZUL, "size": 10.5}), (tono, {"size": 10.5})],
            [("Ideal: ", {"bold": True, "color": DORADO, "size": 10.5}), (prod, {"size": 10.5, "bold": True})],
        ], size=10.5, espacio=5)
    caja(s, MX, Inches(5.95), W - 2 * MX, Inches(0.8), fondo=CREMA)
    texto(s, MX + Inches(0.25), Inches(5.95), W - 2 * MX - Inches(0.5), Inches(0.8),
          [[("FILTRO N·O·A   ", {"bold": True, "color": AZUL, "size": 11}), (C.FILTRO_CALIFICACION, {"size": 11.5})]],
          ancla=MSO_ANCHOR.MIDDLE)
    pie(s, n)


def figura_slide(prs, n, seccion, titulo, fig, intro=""):
    s = nueva(prs)
    encabezado(s, seccion, titulo, intro)
    top = Inches(2.15) if intro else Inches(1.75)
    imagen(s, fig, MX, top, W - 2 * MX, H - top - Inches(0.65))
    pie(s, n)
    return s


def rituales(prs, n):
    s = figura_slide(prs, n, "03 · El proceso de venta", "Las cinco etapas de una conversación",
                     "03_conversacion")
    # la figura ocupa la parte superior; abajo, los rituales
    for sh in list(s.shapes):
        if sh.shape_type == 13:  # imagen: reubicar
            esc = (W - 2 * MX) / sh.width
            sh.left, sh.top = MX, Inches(1.8)
            sh.width, sh.height = int(sh.width * esc), int(sh.height * esc)
    texto(s, MX, Inches(4.75), Inches(6), Inches(0.45), "Rituales del equipo", size=18, bold=True,
          color=AZUL, font=TITULOS)
    w = Inches(3.9)
    for i, (frec, rit) in enumerate(C.RITUALES):
        tarjeta(s, MX + i * (w + Inches(0.2)), Inches(5.25), w, Inches(1.5), frec, rit, size_c=12)


def producto_a(prs, n, p):
    proximo = p.get("proximo")
    s = nueva(prs)
    caja(s, 0, 0, W, Inches(1.65), fondo=P.APAGADO if proximo else AZUL)
    caja(s, 0, Inches(1.65), W, Inches(0.07), fondo=DORADO)
    texto(s, MX, Inches(0.28), Inches(8.5), Inches(0.35), p["etiqueta"].upper(), size=11, bold=True, color=DORADO)
    texto(s, MX, Inches(0.6), Inches(8.5), Inches(0.9), p["nombre"], size=30, bold=True, color=BLANCO, font=TITULOS)
    texto(s, Inches(8.8), Inches(0.28), Inches(3.95), Inches(0.35),
          "PRÓXIMO LANZAMIENTO" if proximo else p.get("etiqueta_precio", "PRECIO SUGERIDO"), size=11, bold=True, color=DORADO,
          alinear=PP_ALIGN.RIGHT)
    texto(s, Inches(8.8), Inches(0.6), Inches(3.95), Inches(0.9), p["precio"],
          size=32 if len(p["precio"]) <= 10 else 24, bold=True, ancla=MSO_ANCHOR.MIDDLE,
          color=BLANCO, alinear=PP_ALIGN.RIGHT)
    texto(s, MX, Inches(1.85), W - 2 * MX, Inches(0.4), p["pago"], size=13, italic=True, color=GRIS)
    texto(s, MX, Inches(2.35), W - 2 * MX, Inches(1.1), [
        [("Qué es.  ", {"bold": True, "color": AZUL}), (p["que_es"], {})],
        [("Para quién.  ", {"bold": True, "color": AZUL}), (p["para_quien"], {})],
    ], size=14, espacio=6)
    w = Inches(5.95)
    for i, (tit, items, marca, col) in enumerate((("Qué incluye", p["incluye"], "✓", VERDE),
                                                   ("Reglas del producto", p["reglas"], "■", DORADO))):
        x = MX + i * (w + Inches(0.25))
        caja(s, x, Inches(3.65), w, Inches(3.1), fondo=CLARO)
        caja(s, x, Inches(3.65), w, Inches(0.06), fondo=col)
        texto(s, x + Inches(0.2), Inches(3.8), w - Inches(0.4), Inches(0.4), tit.upper(), size=11, bold=True, color=AZUL)
        texto(s, x + Inches(0.2), Inches(4.2), w - Inches(0.4), Inches(2.5),
              lista(items, marca=marca, color_marca=col, size=12.5), espacio=5)
    pie(s, n)


def producto_b(prs, n, p):
    s = nueva(prs)
    encabezado(s, f"04 · {p['corto']}", "Cómo presentarlo")
    caja(s, MX, Inches(1.75), Inches(7.3), Inches(2.45), fondo=CREMA)
    caja(s, MX, Inches(1.75), Inches(0.08), Inches(2.45), fondo=DORADO)
    texto(s, MX + Inches(0.3), Inches(1.85), Inches(6.8), Inches(0.35), "SPEECH SUGERIDO", size=11, bold=True, color=AZUL)
    texto(s, MX + Inches(0.3), Inches(2.2), Inches(6.85), Inches(1.95), p["speech"], size=14, italic=True,
          color=AZUL_P, font=TITULOS)
    x2 = Inches(8.2)
    caja(s, x2 - Inches(0.1), Inches(1.75), Inches(4.75), Inches(2.45), fondo=P.PREGUNTAS[1])
    caja(s, x2 - Inches(0.1), Inches(1.75), Inches(4.75), Inches(0.07), fondo=P.PREGUNTAS[0])
    texto(s, x2, Inches(1.9), Inches(4.5), Inches(0.35), "PREGUNTAS DE DESCUBRIMIENTO", size=11, bold=True, color=AZUL)
    texto(s, x2, Inches(2.3), Inches(4.5), Inches(1.85),
          lista(p["preguntas"], marca="?", color_marca=P.PREGUNTAS[0], size=12), espacio=5)
    w = Inches(3.9)
    for i, (tit, txt, (col, fondo)) in enumerate((("Frase de cierre", p["cierre"], P.CIERRE),
                                                  ("Objeción típica", p["objecion"], P.OBJECION),
                                                  ("Tip del asesor", p["tip"], P.TIP))):
        tarjeta(s, MX + i * (w + Inches(0.2)), Inches(4.4), w, Inches(2.45), tit, txt, acento=col, fondo=fondo,
                size_c=12)
    pie(s, n)


def objeciones(prs, n0):
    n = n0
    for i in range(0, len(C.OBJECIONES), 2):
        s = nueva(prs)
        encabezado(s, "05 · Manejo de objeciones", f"Objeciones frecuentes ({i // 2 + 1}/4)")
        for j, (obj, val, rep, ava) in enumerate(C.OBJECIONES[i:i + 2]):
            y = Inches(1.8) + j * Inches(2.5)
            caja(s, MX, y, W - 2 * MX, Inches(0.55), fondo=AZUL)
            texto(s, MX + Inches(0.2), y, W - 2 * MX - Inches(0.4), Inches(0.55), f"{i + j + 1}.  {obj}",
                  size=15, bold=True, color=BLANCO, ancla=MSO_ANCHOR.MIDDLE)
            w = Inches(3.95)
            for k, (letra, nombre, txt) in enumerate((("V", "Validar", f"«{val}»"), ("R", "Reposicionar", rep),
                                                       ("A", "Avanzar", f"«{ava}»"))):
                x = MX + k * (w + Inches(0.15))
                caja(s, x, y + Inches(0.62), w, Inches(1.8), fondo=CLARO)
                texto(s, x + Inches(0.15), y + Inches(0.66), w - Inches(0.3), Inches(1.74), [
                    [(f"{letra}  ", {"font": TITULOS, "size": 16, "bold": True, "color": DORADO}),
                     (nombre.upper(), {"size": 10, "bold": True, "color": AZUL})],
                    [(txt, {"size": 11, "italic": letra != "R"})],
                ], espacio=3)
        pie(s, n)
        n += 1
    return n


def precios(prs, n):
    s = nueva(prs)
    encabezado(s, "06 · Precios, reglas y condiciones", "Tabla de precios",
               "Fuente oficial: Ficha Técnica vigente. Precios sugeridos en US$, IVA incluido.")
    tabla(s, MX, Inches(2.2), W - 2 * MX, ["Producto", "Precio", "Forma de pago", "Mantenimiento", "Plazo / condición"],
          [list(f) for f in C.TABLA_PRECIOS], [2.4, 1.6, 3, 1.7, 3.4], size=12.5, alto_fila=Inches(0.55),
          resaltar=lambda f: f[0] == "Jardín Ceremonial")
    caja(s, MX, Inches(6.2), W - 2 * MX, Inches(0.6), fondo=CREMA)
    texto(s, MX + Inches(0.25), Inches(6.2), W - 2 * MX - Inches(0.5), Inches(0.6),
          [[("EL DESCUENTO NO ES LA SALIDA   ", {"bold": True, "color": AZUL, "size": 10.5}),
            (C.DESCUENTO_NO_ES_SALIDA, {"size": 11})]], ancla=MSO_ANCHOR.MIDDLE)
    pie(s, n)


def precios_existentes(prs, n):
    s = nueva(prs)
    encabezado(s, "06 · Precios, reglas y condiciones", "Portafolio existente: columbarios",
               "Descuento de 15% válido en Finados y durante noviembre. Pago único o hasta 120 meses, con intereses.")
    tabla(s, MX, Inches(2.2), W - 2 * MX, ["Producto", "Precio", "Finados y noviembre (−15%)", "Capacidad", "Incluye"],
          [list(f) for f in C.TABLA_PRECIOS_EXISTENTES], [2.0, 3.6, 3.0, 1.4, 2.8], size=12, alto_fila=Inches(0.62))
    caja(s, MX, Inches(6.2), W - 2 * MX, Inches(0.6), fondo=CREMA)
    texto(s, MX + Inches(0.25), Inches(6.2), W - 2 * MX - Inches(0.5), Inches(0.6),
          [[("UBICACIÓN DEFINIDA O A FUTURO   ", {"bold": True, "color": AZUL, "size": 10.5}),
            (C.UBICACION_DEFINIDA_FUTURO, {"size": 11})]], ancla=MSO_ANCHOR.MIDDLE)
    pie(s, n)


def reglas(prs, n):
    s = nueva(prs)
    encabezado(s, "06 · Precios, reglas y condiciones", "Reglas comerciales")
    tabla(s, MX, Inches(1.8), W - 2 * MX, ["Tema", "Regla"], [list(r) for r in C.REGLAS_COMERCIALES],
          [2.3, 9.8], size=13, alto_fila=Inches(0.6))
    pie(s, n)


def comisiones(prs, n):
    s = nueva(prs)
    encabezado(s, "07 · Política de comisiones", "Comisión única de 8%",
               "Sobre el precio sin IVA. Confirmar con la Gerencia Comercial antes de aplicar.")
    imagen(s, "05_comisiones", MX, Inches(2.1), Inches(6.6), Inches(4.4))
    x = Inches(7.45)
    texto(s, x, Inches(2.15), Inches(5.3), Inches(0.4), "Reglas de comisión", size=18, bold=True,
          color=AZUL, font=TITULOS)
    texto(s, x, Inches(2.65), Inches(5.3), Inches(3.0), lista(C.REGLAS_COMISION, size=12), espacio=5)
    caja(s, x, Inches(5.65), Inches(5.3), Inches(1.1), fondo=CREMA)
    texto(s, x + Inches(0.15), Inches(5.65), Inches(5.0), Inches(1.1), C.COMISION_CUIDADO, size=11.5,
          italic=True, ancla=MSO_ANCHOR.MIDDLE)
    pie(s, n)


def comisiones_tabla(prs, n):
    s = nueva(prs)
    encabezado(s, "07 · Política de comisiones", "Cuándo se paga y cuándo se descuenta")
    tabla(s, MX, Inches(1.8), W - 2 * MX, ["Producto", "%", "Comisión (sin IVA)", "Cuándo se paga", "Estorno"],
          [list(f) for f in C.COMISIONES], [2.6, 0.8, 1.8, 4.6, 2.6], size=12, alto_fila=Inches(0.6),
          resaltar=lambda f: f[0] == "Jardín Ceremonial")
    pie(s, n)


def herramientas(prs, n):
    s = nueva(prs)
    encabezado(s, "08 · Herramientas de apoyo", "Lo que usted tiene a mano")
    for i, (tit, txt) in enumerate(C.HERRAMIENTAS):
        y = Inches(1.8) + i * Inches(0.98)
        caja(s, MX, y, Inches(6.2), Inches(0.85), fondo=CLARO)
        caja(s, MX, y, Inches(0.07), Inches(0.85), fondo=DORADO)
        texto(s, MX + Inches(0.2), y + Inches(0.05), Inches(5.9), Inches(0.8),
              [[(tit, {"bold": True, "color": AZUL, "size": 13})], [(txt, {"size": 11})]], espacio=1)
    x = Inches(7.35)
    caja(s, x, Inches(1.8), Inches(5.4), Inches(4.85), fondo=AZUL)
    texto(s, x + Inches(0.3), Inches(1.95), Inches(4.8), Inches(0.5), "Checklist antes de cerrar", size=18,
          bold=True, color=DORADO, font=TITULOS)
    texto(s, x + Inches(0.3), Inches(2.55), Inches(4.85), Inches(4.0),
          [[("☐  ", {"color": DORADO, "size": 14}), (it, {"color": BLANCO, "size": 13})] for it in C.CHECKLIST_CIERRE],
          espacio=9)
    pie(s, n)


def cierre(prs):
    s = nueva(prs, AZUL)
    caja(s, 0, H - Inches(0.25), W, Inches(0.25), fondo=DORADO)
    texto(s, Inches(1.5), Inches(1.7), W - Inches(3), Inches(2.8), f"«{C.CIERRE_FRASE}»", size=30, italic=True,
          color=BLANCO, font=TITULOS, alinear=PP_ALIGN.CENTER, ancla=MSO_ANCHOR.MIDDLE)
    caja(s, (W - Inches(1.4)) // 2, Inches(4.7), Inches(1.4), Inches(0.06), fondo=DORADO)
    texto(s, Inches(1.5), Inches(5.0), W - Inches(3), Inches(0.5), C.MARCAS, size=15, bold=True,
          color=DORADO, alinear=PP_ALIGN.CENTER)


def main():
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    n = 2
    portada(prs)
    como_usar(prs, n); n += 1
    separador(prs, n, "01", "Quiénes somos y lo que vendemos", C.LEMA); n += 1
    quienes(prs, n); n += 1
    portafolio(prs, n); n += 1
    momentos(prs, n); n += 1
    separador(prs, n, "02", "A quién le vendemos", "Conocer al cliente es la mitad de la venta."); n += 1
    perfiles(prs, n); n += 1
    separador(prs, n, "03", "El proceso de venta", "Un camino claro, del primer contacto a la posventa."); n += 1
    figura_slide(prs, n, "03 · El proceso de venta", "El embudo de prenecesidad", "02_embudo"); n += 1
    rituales(prs, n); n += 1
    separador(prs, n, "04", "Producto por producto", "Qué es, qué incluye, sus reglas y cómo presentarlo."); n += 1
    for p in C.PRODUCTOS:
        producto_a(prs, n, p); n += 1
        if not p.get("proximo"):
            producto_b(prs, n, p); n += 1
        if p["corto"] == "Jardín Ecológico":
            figura_slide(prs, n, "04 · Jardín Ecológico y de Mascotas", "Cómo explicar el ciclo de 3 años",
                         "06_ciclo_ecologico"); n += 1
    separador(prs, n, "04", "Portafolio existente: columbarios", C.INTRO_EXISTENTES); n += 1
    for p in C.PRODUCTOS_EXISTENTES:
        producto_a(prs, n, p); n += 1
        producto_b(prs, n, p); n += 1
    separador(prs, n, "05", "Manejo de objeciones", "Una objeción no es un «no»: es una duda que aún no resolvimos."); n += 1
    figura_slide(prs, n, "05 · Manejo de objeciones", "El método V·R·A", "04_objeciones"); n += 1
    n = objeciones(prs, n)
    separador(prs, n, "06", "Precios, reglas y condiciones", "Transparencia desde la primera conversación."); n += 1
    precios(prs, n); n += 1
    precios_existentes(prs, n); n += 1
    reglas(prs, n); n += 1
    separador(prs, n, "07", "Política de comisiones", "La comisión premia la venta buena, no solo la venta."); n += 1
    comisiones(prs, n); n += 1
    comisiones_tabla(prs, n); n += 1
    separador(prs, n, "08", "Herramientas de apoyo", "Todo lo que necesita para vender con confianza."); n += 1
    herramientas(prs, n); n += 1
    cierre(prs)
    SALIDA.parent.mkdir(exist_ok=True)
    prs.save(SALIDA)
    print(f"✔ {SALIDA} ({len(prs.slides)} diapositivas, {SALIDA.stat().st_size:,} bytes)")


if __name__ == "__main__":
    main()
